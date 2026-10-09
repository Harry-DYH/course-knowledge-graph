"""Build the evidence-based six-page report with the bundled workspace Python."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from create_input_docx import set_font

BASE = Path(__file__).resolve().parents[1]
REPORT = BASE / 'report'
NAME = '真实课程知识抽取测评报告_20261009'


def load(path):
    return json.loads((BASE / path).read_text(encoding='utf-8'))


def p(doc, text, bold=False, style=None):
    para = doc.add_paragraph(style=style)
    para.add_run(text).bold = bold
    return para


def h(doc, title, level=1):
    return doc.add_paragraph(title, style=f'Heading {level}')


def new_page(doc, title):
    para = h(doc, title)
    para.paragraph_format.page_break_before = True


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, title in enumerate(headers):
        t.rows[0].cells[i].text = title
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement('w:cantSplit')
        trpr.append(cannot_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement('w:' + edge)
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')
                border.set(qn('w:color'), 'D9D9D9')
                borders.append(border)
            props.append(borders)
            margin = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), '85' if edge in ('top', 'bottom') else '95')
                item.set(qn('w:type'), 'dxa')
                margin.append(item)
            props.append(margin)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DDE7ED' if ri == 0 else ('F6F8FA' if ri % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = Pt(15)
                para.paragraph_format.keep_with_next = False
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    return t


def main():
    REPORT.mkdir(parents=True, exist_ok=True)
    metadata = {f: load(f'runs/{f}/metadata.json') for f in ('txt', 'docx')}
    metrics = {f: load(f'runs/{f}/scoring/metrics.json') for f in ('txt', 'docx')}
    coverage = {r['format']: r for r in load('diagnostics/chunk_coverage.json')['runs']}
    assert all(not m['completed'] and m['http_status'] == 502 for m in metadata.values())
    assert all(m['raw_counts']['reference_nodes'] == 70 and m['raw_counts']['reference_edges'] == 101 for m in metrics.values())
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(.7)
    sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.78)
    sec.right_margin = Inches(.78)
    sec.footer_distance = Inches(.3)
    set_font(doc.styles['Normal'], 'Songti SC', 10.5)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = Pt(17)
    normal.space_after = Pt(7)
    normal.keep_together = True
    normal.widow_control = True
    set_font(doc.styles['Title'], 'Heiti SC', 24, True)
    doc.styles['Title'].paragraph_format.line_spacing = Pt(30)
    doc.styles['Title'].paragraph_format.space_after = Pt(10)
    for name, size in [('Heading 1', 17), ('Heading 2', 12)]:
        set_font(doc.styles[name], 'Heiti SC', size, True)
        pf = doc.styles[name].paragraph_format
        pf.space_before = Pt(9)
        pf.space_after = Pt(7)
        pf.line_spacing = Pt(22 if size == 17 else 18)
        pf.keep_with_next = True
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    set_font(doc.styles['Footer'], 'Heiti SC', 8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('知识抽取测评  2026年10月9日  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    core = doc.core_properties
    core.title = '真实课程知识抽取测评报告'
    core.subject = 'Hello 算法第7章树 TXT与DOCX真实上传测试和冻结参考匹配评估'
    core.author = '课程知识图谱项目测评'
    core.comments = 'AI辅助测评报告。参考标注与抽样判断尚未获人工确认，不能用于宣称赛题人工准确率验收达标。'

    # Page 1: the outcome and its acceptance limits.
    p(doc, '真实课程知识抽取测评报告', style='Title')
    p(doc, 'Hello 算法 第7章 树', bold=True)
    p(doc, '测评日期 2026年10月9日    项目版本 b510c69')
    h(doc, '本轮结论')
    p(doc, '两次真实文档上传均在内容补全阶段失败，接口返回 HTTP 502。系统保留了部分知识图谱，但本轮没有完成完整章节的抽取闭环，不能判定满足处理成功、60 秒时限或人工准确率 70% 的验收要求。', bold=True)
    p(doc, '本轮以公开中文教材正文替换原有自编小样例，分别上传同内容 TXT 与 DOCX。测试还发现默认分块上限只保留约 47% 的正文 token，7.5 AVL 树和 7.6 小结未进入首轮实体与关系抽取。这是解释漏项的重要诊断证据。[1][2]')
    table(doc, ['观察项', 'TXT', 'DOCX'], [
        ['最终状态', '失败 HTTP 502', '失败 HTTP 502'],
        ['失败前保留的图谱', '50 个知识点\n70 条关系', '47 个知识点\n61 条关系'],
        ['上传至失败响应耗时', '24.553 秒', '24.897 秒'],
        ['正文 token 保留比例', '47.00%', '47.06%'],
    ], [2.5, 2.2, 2.24])
    h(doc, '与原报告及赛题要求的对应', 2)
    p(doc, '原《知识抽取测试与准确率报告》指出，接口跑通与准确率测量应分开，人工真值未建立时不能宣称达标。本报告延续这一口径：给出真实运行证据、冻结参考匹配指标及 AI 辅助抽样复核，不冒充人工验收。[3][4]')
    table(doc, ['赛题要求', '本轮判定'], [
        ['完整课程一章不少于 20 个知识点', '保留数量超过 20，但正文被截断且流程失败；数量不能证明完整章达标。'],
        ['人工抽样准确率不低于 70%', '待人工确认。AI 参考集与 AI 审阅不能替代人工验收。'],
        ['单次解析和抽取总时间不超过 60 秒', '仅测得失败响应时间；不能作为成功完成的耗时成绩。'],
    ], [2.5, 4.44])

    # Page 2: source and reproducibility.
    new_page(doc, '测试资料与运行口径')
    h(doc, '真实资料及授权', 2)
    p(doc, '采用 krahets 及 Hello Algo 贡献者的《Hello 算法》第 7 章树，范围为 7.1 二叉树、7.2 二叉树遍历、7.3 二叉树数组表示、7.4 二叉搜索树、7.5 AVL 树及 7.6 小结，包含原文 Q&A。来源来自官方仓库，获取日期为 2026年10月9日。[1]')
    p(doc, '来源版本固定为 28c1e74c1d3594fca7cc41b78a6bd18b700c5ce9。许可为 CC BY-NC-SA 4.0，原文摘编和相应改编保留署名、非商业及相同方式共享条件；完整许可与原始文件随证据包保存。')
    p(doc, '输入保留 10,344 个字符的文字正文、公式与表格行，去掉插图、多语言代码、交互链接和练习；没有补写定义或先修关系。TXT 与 DOCX 的 219 条非空正文逐行一致。实际解析器会剥离 DOCX 中 46 个行首列表标记；忽略这些标记与空白后内容一致。')
    h(doc, '运行配置', 2)
    table(doc, ['项目', '记录'], [
        ['被测版本', 'course-knowledge-graph\nb510c6933f037daa5d4bb913010d8e719f714b52'],
        ['模型', 'DeepSeek 官方接口；请求模型名 deepseek-flash（别名）；temperature=0、禁用 thinking 为客户端配置值'],
        ['分块与向量', '每块 1,500 token，重叠 100；all-MiniLM-L6-v2，384 维'],
        ['隔离方式', '每种格式使用全新 Neo4j 5.26.0 容器和数据卷，初始节点数 0'],
        ['次数与重试', '每格式上传 1 次；测试脚本不重跑；保留 SDK 默认最多 2 次单请求重试'],
        ['计时边界', '从上传请求发出至 HTTP 响应；包含解析、抽取、入库和补全；不含启动及预热'],
    ], [1.45, 5.49])
    h(doc, '保证结果可复核', 2)
    p(doc, '参考集在读取模型输出之前独立建立并锁定，包含 70 个知识点和 101 条关系，每项附原文证据；它由 AI 辅助编写，尚未人工确认。输入、参考集、协议和输出图谱均保存校验值。没有依据预测结果补加别名或回改参考集，业务源码未修改。')
    p(doc, '两种格式各运行一次，结果同时受到解析分块和模型生成波动影响。本报告不把它们的差异解释为文件格式的因果效果，也不推断多次运行的稳定性。')

    # Page 3: strict matching metrics.
    new_page(doc, '冻结参考集的严格匹配结果')
    p(doc, '以下指标只衡量失败后保留图谱与冻结参考集的匹配程度。节点依据规范名或事先列出的别名匹配，进行 Unicode NFKC、英文大小写及空白归一化；不作模糊匹配，一个预测项最多匹配一个参考项。')
    p(doc, '关系必须双端点、类型均匹配；CONTAINS 和 PREREQUISITE_OF 区分方向，RELATED_TO 不区分方向。重复预测计入未匹配项。精确率 = 匹配数 ÷ 预测数；召回率 = 匹配数 ÷ 参考数；F1 为两者的调和平均数。')
    rows=[]
    for fmt in ('txt', 'docx'):
        for key, label in [('nodes', '知识点'), ('edges', '关系')]:
            m=metrics[fmt][key]
            rows.append([fmt.upper()+' '+label, f"{m['tp']}/{m['prediction_count']}/{m['reference_count']}", f"{m['precision']:.2%}", f"{m['recall']:.2%}", f"{m['f1']:.2%}"])
    table(doc, ['对象', '匹配 预测 参考', '精确率', '召回率', 'F1'], rows, [1.42,1.95,1.19,1.19,1.19])
    p(doc, '两轮各有 35 个参考知识点获得匹配，另有 35 个参考知识点未匹配。TXT 有 15 个预测知识点未匹配，DOCX 有 12 个；其中分别有 2 个、1 个重复预测。关系分别有 53 条、45 条未匹配预测，各包含 1 条重复预测。')
    h(doc, '关系类型明细', 2)
    type_rows=[]
    for fmt in ('txt','docx'):
        for key in ('CONTAINS','RELATED_TO','PREREQUISITE_OF'):
            m=metrics[fmt]['by_edge_type'][key]
            type_rows.append([fmt.upper(), key, f"{m['tp']}/{m['prediction_count']}/{m['reference_count']}",f"{m['precision']:.2%}",f"{m['recall']:.2%}"])
    table(doc, ['格式', '关系类型', '匹配 预测 参考', '精确率', '召回率'],type_rows,[.58,2.05,1.69,1.31,1.31])
    h(doc, '如何解释未匹配项', 2)
    p(doc, '未匹配不等于事实错误。名称粒度、括号中的英文名、树高度与节点高度的区分、操作所属对象，以及参考集未覆盖的合理表述，都可能导致严格匹配失败。错误判断需要回到原文逐项核查。')
    p(doc, '本轮先修关系参考仅 1 条，样本过少，类型指标不能代表整体先修关系识别能力。正文被截断还使全文召回率混合了资料覆盖与抽取质量问题；本轮没有事后缩小参考范围以提高分数。')

    # Page 4: two different failure/coverage findings.
    new_page(doc, '流程失败与正文覆盖诊断')
    h(doc, '内容补全阶段失败', 2)
    p(doc, '两轮均经历抽取、入库并进入内容补全阶段，随后返回 HTTP 502。接口报错为“内容补全未通过原文校验或模型请求失败，可重试；未覆盖已有内容”。这条报错合并了多种原因，现有日志未暴露底层异常，因此不能断定是哪个校验条件或哪次模型请求失败。')
    table(doc, ['阶段首次被观察到', 'TXT', 'DOCX'], [
        ['正在抽取 Extracting', '1.136 秒', '1.146 秒'],
        ['正在入库 Importing', '8.279 秒', '9.267 秒'],
        ['正在补全 Enriching', '9.293 秒', '10.283 秒'],
        ['上传返回 HTTP 502', '24.553 秒', '24.897 秒'],
    ], [3.1,1.92,1.92])
    p(doc, '阶段状态按约 1 秒间隔采样，表中是首次观察时间，不是精确的阶段耗时。DOCX 的末次轮询未采到 Failed，但失败的 HTTP 响应与任务记录均已留存。')
    h(doc, '默认分块上限丢弃后半章', 2)
    p(doc, '原日志与运行后的离线重演均显示：全文产生 13 个文本块，程序仅保留前 6 块，丢弃后 7 块。默认 MAX_TOKEN_CHUNK_SIZE 为 10,000，再除以每块 1,500 并向下取整，得到最多 6 块。考虑块间重叠后，实际保留去重后的前 8,500 token。[2]')
    table(doc, ['覆盖统计', 'TXT', 'DOCX'], [
        ['全文 token 数', f"{coverage['txt']['full_text_tokens']:,}",f"{coverage['docx']['full_text_tokens']:,}"],
        ['保留的去重前缀 token', '8,500', '8,500'],
        ['正文 token 覆盖比例', f"{coverage['txt']['prefix_token_coverage']:.2%}",f"{coverage['docx']['prefix_token_coverage']:.2%}"],
        ['截断位置', '7.4 删除节点开头', '7.4 删除节点开头'],
    ], [3.1,1.92,1.92])
    p(doc, '7.5 AVL 树及 7.6 小结全部未进入首轮实体与关系抽取。这里的覆盖率是去重 token 前缀占全文 token 的比例，不是 6/13 的块数比例，也不是知识点召回率。分块诊断是在运行后独立追加的分析，未改变冻结协议或主评分。')

    # Page 5: blind independent reviews and human work remaining.
    new_page(doc, '原文抽样复核与人工验收边界')
    p(doc, '在每轮保留图谱中按固定种子 20261009 无放回随机抽取 20 个知识点、20 条关系。两名独立 AI 审阅者分别对照原文给出“支持、不支持、不确定”及证据。原文引用已逐字核验；本步骤仍是 AI 辅助审阅，不属于赛题要求的人工抽样验证。')
    p(doc, '汇总规则在主指标之外单独报告：两者均判支持，记为一致支持；两者均判不支持，记为一致不支持；所有其他组合均列为待人工确认，不再由第三个 AI 裁决。')
    table(doc, ['抽样对象', '样本数', '一致支持', '一致不支持', '待人工确认'], [
        ['TXT 知识点',20,20,0,0],['TXT 关系',20,16,2,2],
        ['DOCX 知识点',20,19,0,1],['DOCX 关系',20,18,2,0],
    ], [1.58,.85,1.35,1.5,1.66])
    p(doc, '关系样本的一致支持比例分别为 80% 和 90%。它们只描述已抽出的 20 条关系样本，不衡量未抽出内容，也不能替代全文严格匹配指标或人工准确率。两格式各 40 项中，A/B 判定相同的都是 39 项；相同包括双方均不确定，不能据此证明判断正确。')
    h(doc, '已发现的典型关系问题', 2)
    table(doc, ['输出关系', '两名审阅者共同指出的问题'], [
        ['TXT\n二叉树 CONTAINS 节点深度', '原文把深度定义为节点属性；不支持把它解释为树的组成部分。'],
        ['DOCX\n数组表示 CONTAINS 完美二叉树', '原文用完美二叉树说明数组表示；示例对象并非表示方法的组成部分。'],
        ['DOCX\n完美二叉树 CONTAINS 映射公式', '映射公式表达索引关系；不是完美二叉树包含的结构部件。'],
    ], [3.0,3.94])
    h(doc, '优先交由人工确认的三项', 2)
    p(doc, 'TXT 关系 E03 的两名判断有分歧；TXT 关系 E10 的双方都不确定；DOCX 知识点 N06 涉及“基本操作分组”的粒度判断。清单保留两名审阅者的原始判定和引用，供课程熟悉者填写最终结论。')
    p(doc, '全体 80 个样本仍应由人工复核；不能只检查上述三项后，把其他 AI 判断直接计为人工正确。70 个知识点、101 条关系的参考集也需要人工确认后，才适合作为正式真值集。')

    # Page 6: usable handover, prioritized work, exact sources.
    new_page(doc, '改进顺序与测评交接')
    h(doc, '建议按顺序处理', 2)
    table(doc, ['优先级', '工作与复测证据'], [
        ['1 全文覆盖', '处理分块上限与截断提示，确保完整章节进入抽取；记录生成块数、实际处理块数和正文覆盖率。'],
        ['2 补全失败', '记录可定位的底层错误与失败节点，在保留已有内容的基础上验证恢复策略；完整成功后重新计时。'],
        ['3 关系语义', '区分组成、属性、示例和实现关系；统一 CONTAINS 及先修关系的标注与抽取口径。'],
        ['4 人工确认', '由课程熟悉者复核参考集与样本，明确赛题准确率公式；随后新增批次、重复运行并报告波动。'],
    ], [1.5,5.44])
    p(doc, '以上为后续修复建议。本轮仅建立测评材料、隔离运行与保存证据，没有修改业务源码，也没有通过修改参数后挑选更好结果。正常开发服务与前端仍可使用。')
    h(doc, '可直接接手的证据', 2)
    p(doc, '项目根目录 evaluation/2026-10-09-hello-tree/ 保存本轮证据。下列两份清单是开展人工确认的入口：')
    p(doc, 'review/人工复核清单.md：80 条随机样本、A/B 原始判定、原文引用及人工填写栏。\nreview/参考标注复核清单.md：70 个知识点、101 条关系、别名、原文证据及人工填写栏。')
    p(doc, 'runs/txt/ 与 runs/docx/ 保存 HTTP 响应、任务状态、原始图谱、计时和逐项评分。protocol.json 与 gold/frozen.json 保存预先锁定的口径和校验值；diagnostics/chunk_coverage.json 保存后验覆盖诊断。')
    p(doc, '重算现有主分数只需运行 scripts/score_graph.py，完全离线、不调用 API。重新抽取应使用 scripts/run_benchmark.py 并创建新批次；它会真实调用 DeepSeek。现有批次保留，供对比追踪。')
    h(doc, '来源与依据', 2)
    refs=[
        '[1] Hello 算法，krahets 及贡献者，第7章树。官方页面：www.hello-algo.com/chapter_tree/；官方仓库：github.com/krahets/hello-algo。固定版本、原文校验值与许可见 source/manifest.json 和 source/LICENSE。',
        '[2] 本轮实际运行证据、冻结评分与覆盖诊断：runs/、gold/、review/summary.json、diagnostics/chunk_coverage.json。',
        '[3] 队伍已有《知识抽取测试与准确率报告》，文件“知识抽取准确率测试报告.docx”，区分功能闭环与人工真值指标。',
        '[4] A10 赛题文档《基于AIGC的课程知识图谱智能构建与学习导航系统》，规定不少于20个知识点、人工抽样准确率不低于70%、单次解析抽取不超过60秒。',
    ]
    for text in refs:
        para=p(doc,text)
        para.paragraph_format.line_spacing=Pt(14)
        para.paragraph_format.space_after=Pt(5)
        for run in para.runs: run.font.size=Pt(9)
    output=REPORT/(NAME+'.docx')
    doc.save(output)
    print(output)


if __name__=='__main__':
    main()
