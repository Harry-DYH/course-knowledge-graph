"""Create a DOCX that preserves every nonempty line of the TXT benchmark input.

Run with the bundled workspace Python. Source text, including plain-text
formula notation and Markdown table rows, is intentionally not transformed:
this file tests parser-format equivalence, not mathematical typesetting.
"""
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / 'inputs' / 'Hello算法_第7章树_正文.txt'
OUTPUT = SOURCE.with_suffix('.docx')


def set_font(style, font, size, bold=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + script), font)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)


def is_subheading(text):
    return (len(text) <= 24 and not text.startswith(('-', '|', '表 ', 'Q：'))
            and not re.match(r'^\d+[.、]', text)
            and not any(mark in text for mark in ('。', '：', '？', '！', '=', '\\', '"')))


def main():
    raw = SOURCE.read_text(encoding='utf-8')
    lines = [line for line in raw.splitlines() if line.strip()]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.85)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    section.footer_distance = Inches(.28)
    set_font(doc.styles['Normal'], 'Songti SC', 11)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = Pt(17.5)
    normal.space_after = Pt(6)
    normal.widow_control = True
    normal.keep_together = True
    set_font(doc.styles['Title'], 'Heiti SC', 23, True)
    doc.styles['Title'].paragraph_format.space_after = Pt(14)
    for name, size in [('Heading 1', 16), ('Heading 2', 12)]:
        set_font(doc.styles[name], 'Heiti SC', size, True)
        pf = doc.styles[name].paragraph_format
        pf.keep_with_next = True
        pf.line_spacing = Pt(22 if name == 'Heading 1' else 19)
        pf.space_before = Pt(12 if name == 'Heading 1' else 7)
        pf.space_after = Pt(6)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)

    for i, line in enumerate(lines):
        if i == 0:
            style = 'Title'
        elif re.match(r'^7\.\d+\s', line):
            style = 'Heading 1'
        elif is_subheading(line):
            style = 'Heading 2'
        else:
            style = 'Normal'
        p = doc.add_paragraph(line, style=style)
        if line.startswith('|'):
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = i + 1 < len(lines) and lines[i + 1].startswith('|')
        if line.startswith('表 '):
            p.paragraph_format.keep_with_next = True
        if line.startswith('Q：'):
            p.runs[0].bold = True
            p.paragraph_format.keep_with_next = True
        if line.startswith('- '):
            p.paragraph_format.left_indent = Inches(.1)
            p.paragraph_format.first_line_indent = Inches(-.1)

    # Keep attribution in document properties only: the project's Unstructured
    # parser ingests visible footer text into the extracted benchmark content.
    core = doc.core_properties
    core.title = 'Hello 算法 第7章 树 正文测评输入'
    core.author = 'krahets 及 Hello 算法贡献者'
    core.subject = '知识抽取测评的 DOCX 格式输入，与对应 TXT 的非空行严格一致'
    core.keywords = 'CC BY-NC-SA 4.0; Hello 算法; 树; 知识抽取'
    core.comments = ('来源 https://github.com/krahets/hello-algo/tree/'
                     '28c1e74c1d3594fca7cc41b78a6bd18b700c5ce9/docs/chapter_tree; '
                     '许可 https://creativecommons.org/licenses/by-nc-sa/4.0/; '
                     '仅改变文档排版，正文保留对应 TXT 原始非空行与顺序。')
    doc.save(OUTPUT)
    extracted = [p.text for p in Document(OUTPUT).paragraphs if p.text.strip()]
    assert lines == extracted, 'DOCX body differs from the source TXT nonempty lines'
    report = {
        'source': str(SOURCE), 'output': str(OUTPUT),
        'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        'output_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        'source_nonempty_lines': len(lines),
        'docx_nonempty_paragraphs': len(extracted),
        'exact_nonempty_line_equality': True,
        'whitespace_normalized_equality': re.sub(r'\s+', '', raw) == re.sub(r'\s+', '', '\n'.join(extracted)),
    }
    qa = BASE / 'qa' / 'input_docx'
    qa.mkdir(parents=True, exist_ok=True)
    (qa / 'content_equivalence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
